"""
Helpers compartidos por las pruebas de `procesos`: crear un Juzgado, un
usuario con su PerfilUsuario (rol + juzgado), y una nómina .xlsx sintética
con las mismas cabeceras que ya usan las pruebas del motor
(motor/tests/test_agrupacion.py) — así un cambio de formato de columna
rompe las dos suites a la vez, en vez de que esta quede desincronizada.
"""

import io

import pandas as pd
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document

from nucleo.models import Juzgado, PerfilUsuario, Rol

COLUMNAS_NOMINA = [
    "Patente", "Digito Patente", "Número de Parte", "Rol", "Reg N",
    "Fecha Infracción", "Hora Infracción", "Lugar de la Infracción",
    "Disposiciones Infringidas", "Nombre Infractor", "Apellido 1", "Apellido 2",
    "Cedula de Identidad", "Digito Verifcador", "Domicilio Infractor", "Comuna",
    "Marca del Vehiculo", "Modelo del Vehiculo",
]

CONTRASENA_PRUEBA = "clave-de-prueba-123"  # nunca se usa fuera de las pruebas


def fila_nomina(pat="AAAA11", dig="1", nro_parte="1001", rol="", reg="R-1",
                 fecha="03/09/2026", hora="08:30", lugar="Autopista km 1",
                 disposicion="TAG no vigente", nombre="Juan", ap1="Perez", ap2="Soto",
                 rut="11111111", dv="1", domicilio="Calle 1", comuna="Maipu",
                 marca="Chevrolet", modelo="Spark"):
    return dict(zip(COLUMNAS_NOMINA, [
        pat, dig, nro_parte, rol, reg, fecha, hora, lugar, disposicion,
        nombre, ap1, ap2, rut, dv, domicilio, comuna, marca, modelo,
    ]))


def construir_nomina_xlsx(filas, nombre_archivo="nomina.xlsx"):
    """Arma un .xlsx real (no un .csv con otra extensión) a partir de una
    lista de dicts como los que devuelve `fila_nomina`, listo para subirlo
    como `SimpleUploadedFile` en una prueba de vista."""
    df = pd.DataFrame(filas)
    buffer = io.BytesIO()
    df.to_excel(buffer, index=False)
    buffer.seek(0)
    return SimpleUploadedFile(
        nombre_archivo, buffer.read(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def crear_juzgado(slug="jpl1", nombre="1er Juzgado de Policía Local de Maipú",
                   dia_desde=1, dia_hasta=10, utm_valor=65000, **kwargs):
    return Juzgado.objects.create(
        slug=slug, nombre=nombre, dia_desde=dia_desde, dia_hasta=dia_hasta,
        utm_valor=utm_valor, **kwargs,
    )


def crear_usuario(username, rol=Rol.OPERADOR, juzgado=None, password=CONTRASENA_PRUEBA):
    """Crea un User de Django + su PerfilUsuario. `juzgado=None` con rol
    Administrador es el caso "ve los 3 juzgados" (nucleo.PerfilUsuario);
    con otro rol, None simplemente deja a ese usuario sin juzgado asignado
    (no puede entrar a nada que dependa de uno)."""
    user = get_user_model().objects.create_user(username, password=password)
    PerfilUsuario.objects.create(usuario=user, rol=rol, juzgado=juzgado)
    return user


def ingresar(client, username, password=CONTRASENA_PRUEBA):
    """Login real por /ingresar/ — NO usar Client.login(): django-axes
    exige que la autenticación pase por una request real (AxesBackend
    necesita el `request` que authenticate() no recibe desde login())."""
    return client.post("/ingresar/", {"username": username, "password": password}, follow=True)


def _docx_bytes(doc):
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.read()


def construir_plantilla_valida(texto_distintivo="PLANTILLA DE PRUEBA", nombre_archivo="plantilla.docx",
                                tipo="CITACION"):
    """Una plantilla .docx mínima pero válida para `tipo`: trae los datos
    esenciales que validar_plantilla() exige (nombre, RUT, patente, la
    tabla de causas con su fila repetida y, según el tipo, la cláusula de
    audiencia / fecha de sentencia / monto de la condena) más un texto
    fácil de reconocer en las pruebas, para probar que una plantilla
    PROPIA de un Juzgado (no la por defecto) efectivamente se usa."""
    doc = Document()
    doc.add_paragraph(texto_distintivo)
    doc.add_paragraph(
        "{{ nombre_infractor }} — RUT {{ rut_infractor }} — {{ patente }} — Carta N° {{ numero_carta }}"
    )
    if tipo in ("CITACION", "CITACION_SENTENCIA"):
        doc.add_paragraph("Se cita {{ clausula_audiencia }}")
    if tipo in ("CITACION_SENTENCIA", "SENTENCIA"):
        doc.add_paragraph("Sentencia del {{ fecha_audiencia }}: condena de {{ condena_monto }}")
    tabla = doc.add_table(rows=3, cols=1)
    tabla.rows[0].cells[0].text = "{%tr for causa in causas %}"
    tabla.rows[1].cells[0].text = "{{ causa.rol }} {{ causa.lugar }}"
    tabla.rows[2].cells[0].text = "{%tr endfor %}"
    return SimpleUploadedFile(
        nombre_archivo, _docx_bytes(doc),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


def construir_plantilla_rota(nombre_archivo="plantilla_rota.docx"):
    """Un .docx con un marcador de fila mal escrito a propósito: abre
    "{%tr for causa in causas %}" pero nunca lo cierra con
    "{%tr endfor %}" — el mismo tipo de error que validar_plantilla()
    (motor/plantilla.py) existe para atrapar antes de guardar."""
    doc = Document()
    tabla = doc.add_table(rows=1, cols=1)
    tabla.rows[0].cells[0].text = "{%tr for causa in causas %}{{ causa.nro_parte }}"
    return SimpleUploadedFile(
        nombre_archivo, _docx_bytes(doc),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
