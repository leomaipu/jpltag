"""
Genera motor/plantillas_defecto/citacion_tag.docx: la plantilla que usa
un Juzgado que todavía no subió la suya propia (ver
motor/plantilla.py — PLANTILLA_DEFECTO).

Es un script, no parte de la app en tiempo de ejecución: se corre a
mano cuando hay que regenerar el .docx por defecto (por eso construye
el documento con python-docx igual que el motor anterior — con los
mismos colores y estilo que tenía la carta armada a mano —, pero
dejando marcadores de docxtpl ({{ }} y {%tr %}) como texto en vez de
los datos reales). El resultado SÍ se commitea al repo (es un binario
chico, ~15 KB, y Django lo necesita en tiempo de ejecución).

Uso:
    python scripts/construir_plantilla_defecto.py

Ver motor/plantilla.py para la lista completa de marcadores disponibles.
"""

import os

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

C_HEADER = "1F3864"
C_ROW_PAR = "F2F2F2"
C_NUMERO_CARTA = "DDEBF7"

RUTA_SALIDA = os.path.join(
    os.path.dirname(__file__), "..", "motor", "plantillas_defecto", "citacion_tag.docx",
)


# ══════════════════════════════════════════════════════════════════════
# Helpers de formato docx (solo para este script — construir_contexto()
# y renderizar_carta() en motor/plantilla.py no necesitan nada de esto,
# el documento ya queda formateado acá adentro del .docx)
# ══════════════════════════════════════════════════════════════════════

def set_bg(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)


def set_cell_margins(cell, top=5, bottom=5, left=45, right=45):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for lado, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        nodo = OxmlElement(f"w:{lado}")
        nodo.set(qn("w:w"), str(val))
        nodo.set(qn("w:type"), "dxa")
        tcMar.append(nodo)
    tcPr.append(tcMar)


def add_run(para, texto, bold=False, italic=False, size=10, color="000000", font_family="Arial"):
    r = para.add_run(texto)
    r.font.name = font_family
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = RGBColor.from_string(color)
    return r


def set_para(para, align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6):
    para.alignment = align
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)


def add_border_bottom(paragraph, size=6, color="000000"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def quitar_bordes_tabla(tabla):
    tbl = tabla._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for tag in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    tblPr.append(borders)


# ══════════════════════════════════════════════════════════════════════
# Construcción del documento
# ══════════════════════════════════════════════════════════════════════

def construir():
    doc = Document()
    sec = doc.sections[0]
    m = 1.2 / 2.54
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(m)

    # Encabezado: logo a la izquierda, sello a la derecha (ambos
    # opcionales — si el Juzgado no cargó la imagen, el marcador queda
    # vacío y la celda se ve en blanco, sin romper el layout).
    tabla_enc = doc.add_table(rows=1, cols=2)
    tabla_enc.autofit = True
    c_izq, c_der = tabla_enc.rows[0].cells
    p = c_izq.paragraphs[0]
    set_para(p, WD_ALIGN_PARAGRAPH.LEFT, after=0)
    add_run(p, "{{ logo_img }}", size=10)
    p = c_der.paragraphs[0]
    set_para(p, WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    add_run(p, "{{ sello_img }}", size=10)
    quitar_bordes_tabla(tabla_enc)

    p = doc.add_paragraph()
    set_para(p, WD_ALIGN_PARAGRAPH.CENTER, after=0)
    add_run(p, "{{ juzgado_nombre }}", bold=True, size=11)
    p = doc.add_paragraph()
    set_para(p, WD_ALIGN_PARAGRAPH.CENTER, after=3)
    add_run(p, "{{ ciudad }}", bold=True, size=11)

    p = doc.add_paragraph()
    set_para(p, after=2)
    add_run(p, "{{ ciudad }}, {{ fecha_hoy }}", size=10)
    p = doc.add_paragraph()
    set_para(p, after=2)
    add_run(p, "VISTOS:", bold=True, size=10)

    p = doc.add_paragraph()
    set_para(p, WD_ALIGN_PARAGRAPH.JUSTIFY, after=2)
    add_run(p, "{{ texto_citacion }}", size=10)

    # Tabla de causas: fila de encabezado + 3 filas para que docxtpl
    # repita la fila de datos por cada causa (ver motor/plantilla.py) +
    # la fila final de N° de Carta.
    headers = ["N° Parte", "Rol", "Fecha", "Hora", "Lugar de la Infracción", "Disposición", "Monto"]
    anchos = [Inches(0.55), Inches(0.70), Inches(0.55), Inches(0.38), Inches(2.05), Inches(2.13), Inches(0.85)]

    tabla = doc.add_table(rows=4, cols=7, style="Table Grid")
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    tabla.autofit = False
    tabla.allow_autofit = False

    hdr = tabla.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].width = anchos[i]
        hdr[i].text = ""
        ph = hdr[i].paragraphs[0]
        set_para(ph, WD_ALIGN_PARAGRAPH.CENTER, after=0)
        add_run(ph, h, bold=True, size=7.5, color="FFFFFF")
        set_bg(hdr[i], C_HEADER)
        set_cell_margins(hdr[i])
        hdr[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    fila_for = tabla.rows[1].cells
    for c in fila_for:
        c.width = fila_for[0].width
    celda_for = fila_for[0]
    for c in fila_for[1:]:
        celda_for = celda_for.merge(c)
    celda_for.text = "{%tr for causa in causas %}"
    p = celda_for.paragraphs[0]
    set_para(p, WD_ALIGN_PARAGRAPH.LEFT, after=0)

    alns = [WD_ALIGN_PARAGRAPH.CENTER] * 4 + [WD_ALIGN_PARAGRAPH.LEFT] * 2 + [WD_ALIGN_PARAGRAPH.CENTER]
    marcadores = [
        "{{ causa.nro_parte }}", "{{ causa.rol }}", "{{ causa.fecha }}", "{{ causa.hora }}",
        "{{ causa.lugar }}", "{{ causa.disposicion }}", "{{ causa.monto_txt }}",
    ]
    fila_datos = tabla.rows[2].cells
    for i in range(7):
        fila_datos[i].width = anchos[i]
        fila_datos[i].text = ""
        pc = fila_datos[i].paragraphs[0]
        set_para(pc, alns[i], after=0)
        pc.paragraph_format.line_spacing = 1.0
        add_run(pc, marcadores[i], size=6)
        set_bg(fila_datos[i], C_ROW_PAR)
        set_cell_margins(fila_datos[i])
        fila_datos[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    fila_endfor = tabla.rows[3].cells
    celda_endfor = fila_endfor[0]
    for c in fila_endfor[1:]:
        celda_endfor = celda_endfor.merge(c)
    celda_endfor.text = "{%tr endfor %}"
    p = celda_endfor.paragraphs[0]
    set_para(p, WD_ALIGN_PARAGRAPH.LEFT, after=0)

    fila_numero = tabla.add_row().cells
    celda_numero = fila_numero[0]
    for c in fila_numero[1:]:
        celda_numero = celda_numero.merge(c)
    celda_numero.text = ""
    p = celda_numero.paragraphs[0]
    set_para(p, WD_ALIGN_PARAGRAPH.CENTER, after=0)
    add_run(p, "N° de Carta: {{ numero_carta }}", bold=True, size=8)
    set_bg(celda_numero, C_NUMERO_CARTA)
    set_cell_margins(celda_numero, top=6, bottom=6)
    celda_numero.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    p = doc.add_paragraph()
    set_para(p, after=3)
    add_run(p, "Monto total infracciones: ", bold=True, size=9)
    add_run(p, "{{ monto_total_txt }}", size=9)

    # Firma (imagen si existe, o línea punteada + nombre_firmante — ver
    # LINEA_FIRMA_DEFECTO en motor/plantilla.py, que ya resuelve ese
    # "o" antes de renderizar, así que acá siempre va el mismo marcador).
    p = doc.add_paragraph()
    set_para(p, WD_ALIGN_PARAGRAPH.RIGHT, before=4, after=0)
    add_run(p, "{{ firma_img }}", size=10)
    p = doc.add_paragraph()
    set_para(p, WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    add_run(p, "{{ nombre_firmante }}", bold=True, size=10)
    p = doc.add_paragraph()
    set_para(p, WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    add_run(p, "{{ cargo_firma }}", size=9)

    p = doc.add_paragraph()
    set_para(p, WD_ALIGN_PARAGRAPH.RIGHT, after=3)
    add_run(p, "{{ ciudad }}, {{ fecha_hoy }}", size=10)

    p = doc.add_paragraph()
    set_para(p, after=1)
    add_run(p, "SALIDA DE CARTA", bold=True, size=9)

    p = doc.add_paragraph()
    set_para(p, before=1, after=2)
    add_run(p, "RIJA  Ley 19.841", bold=True, size=16)

    p = doc.add_paragraph()
    set_para(p, after=2)
    add_border_bottom(p)

    # Bloque postal (franqueo) + datos del destinatario.
    tabla_sobre = doc.add_table(rows=1, cols=2)
    tabla_sobre.autofit = True
    c1, c2 = tabla_sobre.rows[0].cells
    p1 = c1.paragraphs[0]
    set_para(p1, after=0)
    add_run(p1, "{{ juzgado_nombre }}", bold=True, size=9)
    p1b = c1.add_paragraph()
    set_para(p1b, after=0)
    add_run(p1b, "{{ ciudad }}", bold=True, size=9)
    p1c = c1.add_paragraph()
    set_para(p1c, after=0)
    add_run(p1c, "{{ direccion_juzgado }}", size=9)

    p2 = c2.paragraphs[0]
    set_para(p2, WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    add_run(p2, "{{ bloque_postal }}", bold=True, size=8)
    p2b = c2.add_paragraph()
    set_para(p2b, WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    add_run(p2b, "Reg N° {{ reg }}", size=8)
    p2c = c2.add_paragraph()
    set_para(p2c, WD_ALIGN_PARAGRAPH.RIGHT, after=0)
    add_run(p2c, "TAG", bold=True, size=8)
    quitar_bordes_tabla(tabla_sobre)

    for label, marcador in [("NOMBRE    : ", "{{ nombre_infractor }}"),
                             ("DIRECCIÓN : ", "{{ domicilio_infractor }}"),
                             ("COMUNA    : ", "{{ comuna_infractor }}")]:
        p = doc.add_paragraph()
        set_para(p, after=1)
        add_run(p, label, bold=True, size=9)
        add_run(p, marcador, size=9)

    return doc


if __name__ == "__main__":
    documento = construir()
    os.makedirs(os.path.dirname(RUTA_SALIDA), exist_ok=True)
    documento.save(RUTA_SALIDA)
    print(f"Guardado: {RUTA_SALIDA}")
