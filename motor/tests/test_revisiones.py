"""
Plantillas con control de cambios y comentarios: las cartas deben salir como
si todos los cambios estuvieran aceptados y sin ningún comentario (Word los
imprime en globos al margen si quedan en el archivo).
"""

import io
import os
import re
import tempfile
import unittest
import zipfile

from docx import Document

from motor.plantilla import (
    TIPO_CITACION, TIPOS_CARTA, renderizar_carta, ruta_plantilla_defecto, texto_de, validar_plantilla,
)
from motor.revisiones import limpiar_docx, plantilla_limpia

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
BASE = ruta_plantilla_defecto("jpl1", "", TIPO_CITACION)

PARRAFO_CON_REVISIONES = (
    f'<w:p {W}>'
    '<w:commentRangeStart w:id="0"/>'
    '<w:r><w:t xml:space="preserve">Texto vigente. </w:t></w:r>'
    '<w:ins w:id="10" w:author="Revisora" w:date="2026-10-01T10:00:00Z">'
    '<w:r><w:t>TEXTO INSERTADO. </w:t></w:r></w:ins>'
    '<w:del w:id="11" w:author="Revisora" w:date="2026-10-01T10:00:00Z">'
    '<w:r><w:delText>TEXTO BORRADO. </w:delText></w:r></w:del>'
    '<w:commentRangeEnd w:id="0"/>'
    '<w:r><w:commentReference w:id="0"/></w:r>'
    '<w:r><w:rPr><w:b/><w:rPrChange w:id="12" w:author="Revisora" w:date="2026-10-01T10:00:00Z"><w:rPr/></w:rPrChange></w:rPr>'
    '<w:t>Final.</w:t></w:r>'
    '</w:p>'
)

COMENTARIOS = (
    f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:comments {W}>'
    '<w:comment w:id="0" w:author="Revisora" w:initials="R"><w:p><w:r><w:t>COMENTARIO INTERNO</w:t></w:r></w:p></w:comment>'
    '</w:comments>'
)


def plantilla_con_revisiones():
    """La plantilla por defecto + un párrafo con cambios, comentario y revisor."""
    with zipfile.ZipFile(BASE) as z:
        partes = {n: z.read(n) for n in z.namelist()}
    doc = partes["word/document.xml"].decode("utf-8")
    doc = re.sub(r"(<w:body>)", lambda m: m.group(1) + PARRAFO_CON_REVISIONES.replace(f" {W}", ""), doc, count=1)
    partes["word/document.xml"] = doc.encode("utf-8")
    partes["word/comments.xml"] = COMENTARIOS.encode("utf-8")
    partes["word/people.xml"] = f'<w:people {W}><w:person w:author="Revisora"/></w:people>'.encode("utf-8")
    rels = partes["word/_rels/document.xml.rels"].decode("utf-8")
    rels = rels.replace(
        "</Relationships>",
        '<Relationship Id="rIdC1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments" Target="comments.xml"/>'
        '<Relationship Id="rIdC2" Type="http://schemas.microsoft.com/office/2011/relationships/people" Target="people.xml"/>'
        "</Relationships>",
    )
    partes["word/_rels/document.xml.rels"] = rels.encode("utf-8")
    tipos = partes["[Content_Types].xml"].decode("utf-8")
    tipos = tipos.replace(
        "</Types>",
        '<Override PartName="/word/comments.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml"/>'
        "</Types>",
    )
    partes["[Content_Types].xml"] = tipos.encode("utf-8")
    ajustes = partes["word/settings.xml"].decode("utf-8")
    ajustes = re.sub(r"(<w:settings[^>]*>)", lambda m: m.group(1) + "<w:trackRevisions/>", ajustes, count=1)
    partes["word/settings.xml"] = ajustes.encode("utf-8")
    salida = io.BytesIO()
    with zipfile.ZipFile(salida, "w", zipfile.ZIP_DEFLATED) as z:
        for nombre, contenido in partes.items():
            z.writestr(nombre, contenido)
    return salida.getvalue()


class LimpiezaDeRevisionesTest(unittest.TestCase):
    def setUp(self):
        self.original = plantilla_con_revisiones()
        self.limpio = limpiar_docx(self.original)

    def _todo_el_xml(self, datos):
        with zipfile.ZipFile(io.BytesIO(datos)) as z:
            return {n: z.read(n).decode("utf-8", "ignore") for n in z.namelist()}

    def test_el_original_de_la_prueba_si_trae_revisiones(self):
        # sanity: si esto falla, la prueba no está probando nada
        xml = self._todo_el_xml(self.original)
        self.assertIn("word/comments.xml", xml)
        self.assertIn("<w:ins ", xml["word/document.xml"])
        self.assertIn("<w:del ", xml["word/document.xml"])

    def test_no_queda_ninguna_marca_de_revision_ni_comentario(self):
        xml = self._todo_el_xml(self.limpio)
        self.assertNotIn("word/comments.xml", xml)
        self.assertNotIn("word/people.xml", xml)
        doc = xml["word/document.xml"]
        for marca in ("<w:ins ", "<w:del ", "w:delText", "commentRange", "commentReference", "rPrChange"):
            self.assertNotIn(marca, doc, marca)
        self.assertNotIn("trackRevisions", xml["word/settings.xml"])
        self.assertNotIn("comments", xml["word/_rels/document.xml.rels"])
        self.assertNotIn("comments", xml["[Content_Types].xml"])
        self.assertNotIn("Revisora", "".join(xml.values()))
        self.assertNotIn("COMENTARIO INTERNO", "".join(xml.values()))

    def test_se_aceptan_los_cambios(self):
        texto = texto_de(Document(io.BytesIO(self.limpio)))
        self.assertIn("Texto vigente. TEXTO INSERTADO. Final.", texto)
        self.assertNotIn("TEXTO BORRADO", texto)

    def test_el_resto_del_documento_queda_igual(self):
        base = texto_de(Document(BASE))
        limpio = texto_de(Document(io.BytesIO(self.limpio)))
        self.assertTrue(limpio.endswith(base))

    def test_una_plantilla_sin_revisiones_no_cambia_de_texto(self):
        self.assertEqual(
            texto_de(Document(io.BytesIO(limpiar_docx(BASE)))), texto_de(Document(BASE)),
        )

    def test_acepta_ruta_bytes_y_archivo_abierto(self):
        with tempfile.TemporaryDirectory() as tmp:
            ruta = os.path.join(tmp, "con_revisiones.docx")
            with open(ruta, "wb") as f:
                f.write(self.original)
            desde_ruta = limpiar_docx(ruta)
            self.assertIn("word/document.xml", self._todo_el_xml(desde_ruta))
            self.assertEqual(plantilla_limpia(ruta).read()[:2], b"PK")
        self.assertEqual(plantilla_limpia(io.BytesIO(self.original)).read()[:2], b"PK")


class CartasSinRevisionesTest(unittest.TestCase):
    def test_la_carta_generada_no_lleva_comentarios_ni_cambios(self):
        from motor.plantilla import contexto_de_ejemplo
        with tempfile.TemporaryDirectory() as tmp:
            ruta = os.path.join(tmp, "plantilla.docx")
            with open(ruta, "wb") as f:
                f.write(plantilla_con_revisiones())
            doc = renderizar_carta(ruta, contexto_de_ejemplo(TIPO_CITACION), {})
            salida = io.BytesIO()
            doc.save(salida)
        with zipfile.ZipFile(io.BytesIO(salida.getvalue())) as z:
            nombres = z.namelist()
            documento = z.read("word/document.xml").decode("utf-8")
        self.assertFalse([n for n in nombres if "comments" in n or "people" in n])
        for marca in ("<w:ins ", "<w:del ", "w:delText", "commentRange", "commentReference"):
            self.assertNotIn(marca, documento, marca)
        self.assertNotIn("TEXTO BORRADO", documento)
        self.assertIn("TEXTO INSERTADO", documento)

    def test_validar_plantilla_acepta_una_plantilla_con_revisiones(self):
        for tipo in TIPOS_CARTA:
            ruta = ruta_plantilla_defecto("jpl2", "", tipo)
            es_valida, mensaje = validar_plantilla(ruta, tipo)
            self.assertTrue(es_valida, mensaje)
        es_valida, mensaje = validar_plantilla(io.BytesIO(plantilla_con_revisiones()), TIPO_CITACION)
        self.assertTrue(es_valida, mensaje)
