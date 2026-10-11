"""
Las 9 plantillas por defecto (3 Juzgados x 3 tipos de carta): todas
renderizan con datos SINTÉTICOS, no dejan marcadores sin reemplazar, y
cada una conserva el texto propio de su Juzgado (ordinal del Juzgado,
diferencias de redacción del 2do). Ninguna prueba usa datos reales.
"""

import io
import os
import tempfile
import unittest

from docx import Document

from motor.cartas import condena_monto_txt, fmt_utm_simple
from motor.plantilla import (
    TIPOS_CARTA,
    PlantillaNoDisponible,
    construir_contexto,
    generar_cartas_certificadas,
    numero_juzgado,
    renderizar_carta,
    resolver_plantilla,
    ruta_plantilla_defecto,
    texto_de,
    validar_plantilla,
)

ORDINAL = {1: "1er", 2: "2do", 3: "3er"}


def _info(montos=(1,)):
    return {
        "digPat": "7", "rol_raw": "", "reg_raw": "R-77",
        "nombre": "MARIA PRUEBA SINTETICA", "rut": "99999999-9",
        "domicilio": "Calle Inventada 123", "comuna": "Maipu",
        "vehiculo": "Marca X, Modelo Y",
        "causas": [
            {"nroParte": str(1000 + i), "rol_raw": "", "fecha": f"0{i + 1}/09/2026",
             "hora": "08:30", "lugar": f"Autopista Sintetica km {i}",
             "disposicion": "TAG no vigente", "monto": m}
            for i, m in enumerate(montos)
        ],
    }


def _config(n, tipo, **extra):
    cfg = {
        "tipo_carta": tipo, "slug_juzgado": f"jpl{n}",
        "nombre_juzgado": f"JUZGADO {n} (PRUEBA)", "utm_valor": 65000,
        "nombre_firmante": "ANA FIRMANTE DE PRUEBA", "cargo_firma": "SECRETARIA ABOGADA",
        "fecha_audiencia": "02/10/2026", "max_causas_por_pagina": 20,  # explícito: no depende del defecto
    }
    cfg.update(extra)
    return cfg


def _pagina(info, parte=1, total_partes=1, chunk=None):
    chunk = chunk if chunk is not None else info["causas"]
    return {"patente": "ZZZZ99", "info": info, "chunk": chunk, "parte": parte,
            "total_partes": total_partes, "total_causas": len(info["causas"])}


_texto = texto_de


def _renderizar(n, tipo, info=None, **extra):
    cfg = _config(n, tipo, **extra)
    pg = _pagina(info or _info((1, 1)))
    ctx = construir_contexto(pg, ("06/10/2026", "09:30"), 1, cfg)
    return _texto(renderizar_carta(resolver_plantilla(cfg, tipo), ctx, cfg))


class TestPlantillasPorDefecto(unittest.TestCase):
    def test_las_nueve_validan(self):
        for n in (1, 2, 3):
            for tipo in TIPOS_CARTA:
                with self.subTest(juzgado=n, tipo=tipo):
                    ok, msg = validar_plantilla(ruta_plantilla_defecto(f"jpl{n}", "", tipo), tipo)
                    self.assertTrue(ok, msg)

    def test_renderizan_sin_marcadores_sueltos_y_con_los_datos(self):
        for n in (1, 2, 3):
            for tipo in TIPOS_CARTA:
                with self.subTest(juzgado=n, tipo=tipo):
                    t = _renderizar(n, tipo)
                    for resto in ("{{", "}}", "{%", "%}"):
                        self.assertNotIn(resto, t)
                    for dato in ("MARIA PRUEBA SINTETICA", "99999999-9", "Calle Inventada 123",
                                 "ZZZZ99-7", "Marca X, Modelo Y", "ANA FIRMANTE DE PRUEBA",
                                 "SECRETARIA ABOGADA", "Autopista Sintetica km 1"):
                        self.assertIn(dato, t)

    def test_cada_juzgado_nombra_su_propio_ordinal(self):
        for n in (1, 2, 3):
            for tipo in TIPOS_CARTA:
                with self.subTest(juzgado=n, tipo=tipo):
                    t = _renderizar(n, tipo)
                    self.assertIn(f"en el {ORDINAL[n]} Juzgado de Policía Local", t)
                    for otro in set(ORDINAL) - {n}:
                        self.assertNotIn(f"en el {ORDINAL[otro]} Juzgado", t)

    def test_citacion_usa_la_agenda_y_no_condena(self):
        for n in (1, 2, 3):
            t = _renderizar(n, "CITACION")
            self.assertIn("el día 06/10/2026, a las 09:30 hrs.", t)
            self.assertNotIn("CONDENADO", t)

    def test_citacion_y_sentencia_trae_la_condena(self):
        for n in (1, 2, 3):
            t = _renderizar(n, "CITACION_SENTENCIA")
            self.assertIn("el día 06/10/2026, a las 09:30 hrs.", t)
            self.assertIn("será CONDENADO(A) al pago de 1 UTM por infracción", t)

    def test_sentencia_usa_la_fecha_de_la_sentencia_y_no_la_agenda(self):
        for n in (1, 2, 3):
            t = _renderizar(n, "SENTENCIA")
            self.assertIn("que con fecha 02/10/2026, atendida la rebeldía, se condena al pago de 1 UTM", t)
            self.assertNotIn("06/10/2026", t)
            self.assertNotIn("09:30", t)

    def test_el_2do_juzgado_conserva_su_redaccion_propia(self):
        """Los ejemplos del 2do JPL no dicen 'a prestar declaración
        indagatoria'; el 1ro y el 3ro sí."""
        self.assertNotIn("a prestar declaración indagatoria", _renderizar(2, "CITACION"))
        for n in (1, 3):
            self.assertIn("a prestar declaración indagatoria", _renderizar(n, "CITACION"))

    def test_sin_agenda_la_citacion_dice_a_coordinar(self):
        cfg = _config(1, "CITACION")
        from motor.agenda import SLOT_SIN_AGENDA
        ctx = construir_contexto(_pagina(_info()), SLOT_SIN_AGENDA, 1, cfg)
        t = _texto(renderizar_carta(resolver_plantilla(cfg, "CITACION"), ctx, cfg))
        self.assertIn("en fecha y hora a coordinar con el Juzgado", t)

    def test_multas_distintas_no_afirman_un_monto_unico(self):
        t = _renderizar(1, "CITACION_SENTENCIA", info=_info((1, 2)))
        self.assertIn("al pago de la multa indicada por infracción", t)
        self.assertIn("2 UTM", t)

    def test_patente_con_mas_causas_que_una_pagina_avisa_la_parte(self):
        info = _info((1,) * 25)
        cfg = _config(1, "CITACION", max_causas_por_pagina=20)
        pg = _pagina(info, parte=2, total_partes=2, chunk=info["causas"][20:])
        ctx = construir_contexto(pg, ("06/10/2026", "09:30"), 1, cfg)
        t = _texto(renderizar_carta(resolver_plantilla(cfg, "CITACION"), ctx, cfg))
        self.assertIn("Parte 2 de 2: infracciones 21 a 25 de 25", t)
        una = _renderizar(1, "CITACION")
        self.assertNotIn("Parte", una)


class TestResolverPlantilla(unittest.TestCase):
    def test_numero_juzgado(self):
        self.assertEqual(numero_juzgado("jpl2"), 2)
        self.assertEqual(numero_juzgado("primer-jpl"), None)
        self.assertEqual(numero_juzgado("x", "TERCER JUZGADO DE POLICÍA LOCAL"), 3)
        self.assertIsNone(numero_juzgado("", ""))

    def test_juzgado_desconocido_falla_en_vez_de_usar_texto_de_otro(self):
        with self.assertRaises(PlantillaNoDisponible):
            resolver_plantilla({"slug_juzgado": "otro", "nombre_juzgado": "Otro"}, "CITACION")
        with self.assertRaises(PlantillaNoDisponible):
            resolver_plantilla({"slug_juzgado": "jpl9"}, "CITACION")

    def test_tipo_desconocido(self):
        with self.assertRaises(PlantillaNoDisponible):
            resolver_plantilla({"slug_juzgado": "jpl1"}, "OTRA")

    def test_la_plantilla_propia_gana_a_la_del_repo(self):
        with tempfile.TemporaryDirectory() as d:
            ruta = os.path.join(d, "propia.docx")
            Document().save(ruta)
            cfg = {"slug_juzgado": "jpl1", "plantillas_word": {"SENTENCIA": ruta}}
            self.assertEqual(resolver_plantilla(cfg, "SENTENCIA"), ruta)
            self.assertEqual(resolver_plantilla(cfg, "CITACION"), ruta_plantilla_defecto("jpl1", "", "CITACION"))

    def test_validar_rechaza_plantilla_incompleta(self):
        with tempfile.TemporaryDirectory() as d:
            ruta = os.path.join(d, "vacia.docx")
            doc = Document()
            doc.add_paragraph("Solo texto, sin ningún dato")
            doc.save(ruta)
            ok, msg = validar_plantilla(ruta, "CITACION")
            self.assertFalse(ok)
            self.assertIn("nombre del infractor", msg)
            self.assertIn("la patente", msg)

    def test_citacion_no_pasa_como_sentencia(self):
        ok, msg = validar_plantilla(ruta_plantilla_defecto("jpl1", "", "CITACION"), "SENTENCIA")
        self.assertFalse(ok)
        self.assertIn("fecha de la sentencia", msg)


class TestHelpersDeTexto(unittest.TestCase):
    def test_fmt_utm_simple(self):
        self.assertEqual(fmt_utm_simple(1), "1 UTM")
        self.assertEqual(fmt_utm_simple(1.5), "1,5 UTM")
        self.assertEqual(fmt_utm_simple(2.0), "2 UTM")

    def test_condena_monto_txt(self):
        self.assertEqual(condena_monto_txt([{"monto": 1}, {"monto": 1}]), "1 UTM")
        self.assertEqual(condena_monto_txt([{"monto": 1}, {"monto": 2}]), "la multa indicada")


class TestGenerarPorTipo(unittest.TestCase):
    def test_genera_el_docx_y_el_seguimiento_con_el_tipo(self):
        import pandas as pd
        with tempfile.TemporaryDirectory() as d:
            cfg = _config(3, "SENTENCIA")
            grupos = {"ZZZZ99": _info((1, 1)), "YYYY88": _info((1,))}
            res = generar_cartas_certificadas(grupos, "jpl3", {}, 0, cfg, carpeta_salida=d)
            self.assertEqual(res["errores"], [])
            doc = Document(res["archivos"][0])
            t = _texto(doc)
            self.assertEqual(t.count("NOTIFÍQUESE la SENTENCIA"), 2)
            df = pd.read_excel(res["seguimiento"])
            self.assertEqual(set(df["Tipo de carta"]), {"SENTENCIA"})
            self.assertEqual(len(df), 3)

    def test_sin_plantilla_disponible_levanta_error_claro(self):
        with tempfile.TemporaryDirectory() as d:
            cfg = _config(1, "CITACION", slug_juzgado="desconocido", nombre_juzgado="Sin número")
            with self.assertRaises(PlantillaNoDisponible):
                generar_cartas_certificadas({"ZZZZ99": _info()}, "x", {}, 0, cfg, carpeta_salida=d)


if __name__ == "__main__":
    unittest.main()
